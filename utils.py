import os
import base64
import re
import io
import time
from email.utils import parsedate_to_datetime

from google import genai
from google.genai import types as genai_types
from PyPDF2 import PdfReader
from docx import Document
from openpyxl import load_workbook
from PIL import Image


# ─── Gemini Setup ──────────────────────────────────────────────────────────────

# Read model from env so it can be overridden; default to gemini-2.5-flash
MODEL_ID = os.environ.get("GEMINI_MODEL", "gemini-2.5-flash")

def get_gemini_client():
    api_key = os.environ.get("GEMINI_API_KEY", "")
    if not api_key:
        raise ValueError("GEMINI_API_KEY not set in environment.")
    return genai.Client(api_key=api_key)


# ─── Attachment Text Extraction ────────────────────────────────────────────────

def extract_text_from_pdf(data: bytes) -> str:
    try:
        reader = PdfReader(io.BytesIO(data))
        texts = []
        for page in reader.pages:
            t = page.extract_text()
            if t:
                texts.append(t)
        return "\n".join(texts) if texts else "[PDF had no extractable text]"
    except Exception as e:
        return f"[Could not read PDF: {e}]"


def extract_text_from_docx(data: bytes) -> str:
    try:
        doc = Document(io.BytesIO(data))
        return "\n".join(p.text for p in doc.paragraphs if p.text.strip()) or "[DOCX was empty]"
    except Exception as e:
        return f"[Could not read DOCX: {e}]"


def extract_text_from_xlsx(data: bytes) -> str:
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        lines = []
        for sheet in wb.worksheets:
            lines.append(f"[Sheet: {sheet.title}]")
            for row in sheet.iter_rows(values_only=True):
                row_text = "\t".join(str(c) if c is not None else "" for c in row)
                if row_text.strip():
                    lines.append(row_text)
        return "\n".join(lines) if lines else "[XLSX was empty]"
    except Exception as e:
        return f"[Could not read XLSX: {e}]"


def extract_text_from_image(data: bytes, filename: str, client) -> str:
    """Send image directly to Gemini for description."""
    try:
        img = Image.open(io.BytesIO(data))
        img_bytes = io.BytesIO()
        fmt = img.format or "PNG"
        img.save(img_bytes, format=fmt)
        img_bytes.seek(0)

        mime_map = {"JPEG": "image/jpeg", "JPG": "image/jpeg", "PNG": "image/png",
                    "GIF": "image/gif", "WEBP": "image/webp", "BMP": "image/bmp"}
        mime = mime_map.get(fmt.upper(), "image/png")

        response = client.models.generate_content(
            model=MODEL_ID,
            contents=[
                f"Describe this image in detail (filename: {filename}):",
                genai_types.Part.from_bytes(data=img_bytes.read(), mime_type=mime),
            ]
        )
        return response.text
    except Exception as e:
        return f"[Could not process image: {e}]"


def extract_attachment_content(part, client) -> dict:
    """Extract content from a single MIME part attachment."""
    filename = part.get_filename() or "unnamed"
    data = part.get_payload(decode=True)
    if not data:
        return {"filename": filename, "content": "[No data]"}

    ct = part.get_content_type().lower()
    name_lower = filename.lower()

    if ct == "application/pdf" or name_lower.endswith(".pdf"):
        content = extract_text_from_pdf(data)
    elif ct in ("application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                "application/msword") or name_lower.endswith((".docx", ".doc")):
        content = extract_text_from_docx(data)
    elif ct in ("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                "application/vnd.ms-excel") or name_lower.endswith((".xlsx", ".xls")):
        content = extract_text_from_xlsx(data)
    elif ct.startswith("image/") or name_lower.endswith((".png", ".jpg", ".jpeg", ".gif", ".webp", ".bmp")):
        content = extract_text_from_image(data, filename, client)
    elif ct.startswith("text/"):
        try:
            content = data.decode("utf-8", errors="replace")
        except Exception:
            content = "[Could not decode text file]"
    else:
        content = f"[Unsupported file type: {ct}]"

    return {"filename": filename, "content": content[:4000]}  # cap per attachment


# ─── Gmail Message Parsing ─────────────────────────────────────────────────────

def parse_gmail_message(msg_data: dict, client) -> dict:
    """Parse a Gmail API message into a structured dict with attachments extracted."""
    headers = {h["name"]: h["value"] for h in msg_data.get("payload", {}).get("headers", [])}

    subject = headers.get("Subject", "(No Subject)")
    sender  = headers.get("From", "Unknown")
    date_str = headers.get("Date", "")
    try:
        date = parsedate_to_datetime(date_str).strftime("%Y-%m-%d %H:%M")
    except Exception:
        date = date_str or "Unknown date"

    body_text = ""
    attachments = []

    def walk_parts(parts):
        nonlocal body_text
        for part in parts:
            ct = part.get("mimeType", "")
            if ct == "text/plain":
                data = part.get("body", {}).get("data", "")
                if data:
                    body_text += base64.urlsafe_b64decode(data + "==").decode("utf-8", errors="replace")
            elif ct == "text/html" and not body_text:
                data = part.get("body", {}).get("data", "")
                if data:
                    html = base64.urlsafe_b64decode(data + "==").decode("utf-8", errors="replace")
                    # strip tags crudely
                    body_text += re.sub(r"<[^>]+>", " ", html)
            elif part.get("filename"):
                att_data = part.get("body", {}).get("data", "")
                if att_data:
                    raw = base64.urlsafe_b64decode(att_data + "==")
                    # Create a fake MIME part for our extractor
                    import email.mime.base
                    fake_part = email.mime.base.MIMEBase(*ct.split("/", 1) if "/" in ct else (ct, ""))
                    fake_part.set_payload(raw)
                    fake_part.add_header("Content-Disposition", "attachment", filename=part["filename"])
                    attachments.append(extract_attachment_content(fake_part, client))
            elif "parts" in part:
                walk_parts(part["parts"])

    payload = msg_data.get("payload", {})
    if "parts" in payload:
        walk_parts(payload["parts"])
    else:
        data = payload.get("body", {}).get("data", "")
        if data:
            body_text = base64.urlsafe_b64decode(data + "==").decode("utf-8", errors="replace")

    body_text = re.sub(r"\s{3,}", "\n\n", body_text).strip()

    return {
        "subject": subject,
        "sender": sender,
        "date": date,
        "body": body_text[:6000],
        "attachments": attachments,
        "id": msg_data.get("id", ""),
    }


# ─── Summarization ─────────────────────────────────────────────────────────────

def summarize_email(email_data: dict, client) -> str:
    """Build a Gemini prompt and return a summary string."""
    att_section = ""
    for att in email_data["attachments"]:
        att_section += f"\n\n--- Attachment: {att['filename']} ---\n{att['content']}"

    prompt = f"""You are an intelligent email assistant. Summarize the following email clearly and concisely.

**From:** {email_data['sender']}
**Date:** {email_data['date']}
**Subject:** {email_data['subject']}

**Body:**
{email_data['body'] or '(No body text)'}
{att_section}

Please provide:
1. A 2-3 sentence **overview** of what this email is about
2. **Key points** (bullet list)
3. **Action items** (if any, else say "None")
4. **Attachments summary** (if any, else say "None")

Format your response with clear headings using markdown."""

    # Retry with exponential backoff for transient 503 errors only.
    # Do NOT retry 429 quota-exceeded errors — they won't clear for hours.
    max_retries = 4
    wait = 2  # seconds
    for attempt in range(max_retries):
        try:
            response = client.models.generate_content(
                model=MODEL_ID,
                contents=prompt
            )
            return response.text
        except Exception as e:
            err = str(e)
            # Quota exhausted — stop immediately, no point retrying
            if "429" in err or "RESOURCE_EXHAUSTED" in err:
                return ("[Quota exceeded] You've hit the daily free-tier limit for the Gemini API. "
                        "Please wait a few hours for the quota to reset, or upgrade your "
                        "Google AI Studio plan at https://ai.google.dev/gemini-api/docs/rate-limits")
            # Retry only on transient server-side 503 errors
            if attempt < max_retries - 1 and any(x in err for x in ["503", "UNAVAILABLE"]):
                time.sleep(wait)
                wait *= 2   # exponential backoff: 2s → 4s → 8s
                continue
            return f"[Summarization failed: {e}]"


# ─── Draft Reply ────────────────────────────────────────────────────────────────

def draft_reply(subject: str, sender: str, summary: str,
                user_instructions: str, client) -> str:
    """Use Gemini to compose a professional email reply."""
    instruction_block = ""
    if user_instructions and user_instructions.strip():
        instruction_block = f"\n\nAdditional instructions from the user:\n{user_instructions.strip()}"

    prompt = f"""You are a professional email assistant. Write a clear, polite, and concise reply to the following email.

**Original Email Details:**
- From: {sender}
- Subject: {subject}
- Summary: {summary}
{instruction_block}

**Instructions:**
- Write ONLY the body of the reply email (no subject line, no "From:", no "To:")
- Start with a polite greeting (e.g. "Hi [Name],")
- Keep it professional and concise
- End with a proper sign-off (e.g. "Best regards,\\n[Your Name]")
- Do NOT include any meta-commentary or explanations — just the email body"""

    max_retries = 4
    wait = 2
    for attempt in range(max_retries):
        try:
            response = client.models.generate_content(
                model=MODEL_ID,
                contents=prompt
            )
            return response.text
        except Exception as e:
            err = str(e)
            if "429" in err or "RESOURCE_EXHAUSTED" in err:
                return "[Quota exceeded] Cannot draft reply — daily API limit reached."
            if attempt < max_retries - 1 and any(x in err for x in ["503", "UNAVAILABLE"]):
                time.sleep(wait)
                wait *= 2
                continue
            return f"[Draft failed: {e}]"

